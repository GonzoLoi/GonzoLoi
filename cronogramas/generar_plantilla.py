#!/usr/bin/env python3
"""Genera un Excel de plantilla (con datos de ejemplo) para completar y luego
pasarle a main.py.

Uso:
    python generar_plantilla.py plantilla.xlsx
    python generar_plantilla.py plantilla.xlsx --duracion 45 --modulos 6 \
        --inicio 08:00 --recreo 2:15 --recreo 4:15
"""

import argparse
import sys

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from scheduler.periods import generar_periodos

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _header_row(ws, headers: list[str]) -> None:
    for col_i, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=col_i, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        ws.column_dimensions[get_column_letter(col_i)].width = max(14, len(h) + 4)


def _parse_recreo(valor: str) -> tuple[int, int]:
    try:
        despues, duracion = valor.split(":")
        return int(despues), int(duracion)
    except ValueError as exc:
        raise argparse.ArgumentTypeError(
            f"--recreo debe tener el formato 'después_de_modulo:duracion_min', recibido '{valor}'"
        ) from exc


def build_template(
    path: str,
    duracion_modulo: int = 45,
    cantidad_modulos: int = 6,
    hora_inicio: str = "08:00",
    recreos: list[tuple[int, int]] | None = None,
) -> None:
    recreos = recreos if recreos is not None else [(2, 15), (4, 15)]
    periodos, recreos_franjas = generar_periodos(hora_inicio, duracion_modulo, cantidad_modulos, recreos)

    wb = Workbook()

    ws = wb.active
    ws.title = "Config"
    _header_row(ws, ["Dias", "Periodos", "Recreos"])
    dias = ["Lunes", "Martes", "Miércoles", "Jueves", "Viernes"]
    max_filas = max(len(dias), len(periodos), len(recreos_franjas))
    for i in range(max_filas):
        ws.cell(row=2 + i, column=1, value=dias[i] if i < len(dias) else None)
        ws.cell(row=2 + i, column=2, value=periodos[i] if i < len(periodos) else None)
        ws.cell(row=2 + i, column=3, value=recreos_franjas[i] if i < len(recreos_franjas) else None)

    ws = wb.create_sheet("Asignaciones")
    _header_row(ws, ["Grupo", "Materia", "Profesor", "HorasSemana", "Aula"])
    ejemplo = [
        ("1ro A", "Matemática", "Juana Pérez", 4, "Aula 1"),
        ("1ro A", "Lengua", "Marcos Ruiz", 4, "Aula 1"),
        ("1ro A", "Inglés", "Laura Gómez", 2, "Aula 1"),
        ("1ro B", "Matemática", "Juana Pérez", 4, "Aula 2"),
        ("1ro B", "Lengua", "Marcos Ruiz", 4, "Aula 2"),
        ("1ro B", "Inglés", "Laura Gómez", 2, "Aula 2"),
    ]
    for row_i, (grupo, materia, profesor, horas, aula) in enumerate(ejemplo, start=2):
        ws.cell(row=row_i, column=1, value=grupo)
        ws.cell(row=row_i, column=2, value=materia)
        ws.cell(row=row_i, column=3, value=profesor)
        ws.cell(row=row_i, column=4, value=horas)
        ws.cell(row=row_i, column=5, value=aula)

    ws = wb.create_sheet("NoDisponibilidad")
    _header_row(ws, ["Profesor", "Dia", "Periodo"])
    ws.cell(row=2, column=1, value="Marcos Ruiz")
    ws.cell(row=2, column=2, value="Viernes")
    ws.cell(row=2, column=3, value=periodos[0])

    notas = wb.create_sheet("Instrucciones")
    notas.sheet_view.showGridLines = False
    lines = [
        "Cómo completar este archivo:",
        "",
        "Hoja 'Config': listá los días de la semana (columna Dias), los",
        "  módulos de clase (columna Periodos) y, si aplica, los recreos",
        "  (columna Recreos). Las tres columnas son independientes, no hace",
        "  falta que tengan la misma cantidad de filas. Esta plantilla ya",
        f"  viene con módulos de {duracion_modulo} minutos empezando a las",
        f"  {hora_inicio} y recreos entre medio — podés regenerarla con otra",
        "  duración usando, por ejemplo:",
        "    python generar_plantilla.py plantilla.xlsx --duracion 60 \\",
        "        --modulos 5 --inicio 08:00 --recreo 2:20",
        "",
        "Hoja 'Asignaciones': una fila por cada combinación Grupo+Materia+",
        "  Profesor, con la cantidad de módulos semanales que necesita esa",
        "  materia (columna HorasSemana) y, opcionalmente, el aula donde se",
        "  dicta. Borrá las filas de ejemplo y cargá las tuyas.",
        "",
        "Hoja 'NoDisponibilidad' (opcional): si un profesor NO puede dar",
        "  clase en un Día+Periodo puntual, agregalo acá. Se puede dejar",
        "  vacía o borrar la hoja si no aplica.",
        "",
        "Los recreos de la hoja Config nunca reciben clases: el programa los",
        "  muestra en el cronograma final como una fila 'RECREO' pero no los",
        "  usa para agendar materias.",
        "",
        "Una vez completado, generá el cronograma con:",
        "  python main.py plantilla.xlsx cronograma.xlsx",
    ]
    for i, line in enumerate(lines, start=1):
        notas.cell(row=i, column=1, value=line)
    notas.column_dimensions["A"].width = 78

    wb.save(path)


def parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("path", nargs="?", default="plantilla.xlsx", help="Archivo .xlsx a generar")
    parser.add_argument(
        "--duracion", type=int, default=45, help="Duración de cada módulo en minutos (típico: 45 o 60)"
    )
    parser.add_argument("--modulos", type=int, default=6, help="Cantidad de módulos de clase por día")
    parser.add_argument("--inicio", default="08:00", help="Hora de inicio del primer módulo, formato HH:MM")
    parser.add_argument(
        "--recreo",
        action="append",
        type=_parse_recreo,
        default=None,
        metavar="despues_de_modulo:duracion_min",
        help="Recreo tras un módulo dado, ej: --recreo 2:15. Repetible para varios recreos.",
    )
    return parser.parse_args(argv)


def main() -> int:
    args = parse_args(sys.argv[1:])
    build_template(args.path, args.duracion, args.modulos, args.inicio, args.recreo)
    print(f"Plantilla generada en '{args.path}'.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
