"""Read the input workbook (Config / Asignaciones / NoDisponibilidad) and write
the resulting weekly schedule as a formatted .xlsx workbook."""

import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .models import Assignment, Config, Slot

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
CELL_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _header_map(ws: Worksheet) -> dict[str, int]:
    headers = {}
    for cell in ws[1]:
        if cell.value is not None:
            headers[str(cell.value).strip().lower()] = cell.column
    return headers


def _col(headers: dict[str, int], *names: str, required: bool = True, sheet: str = "") -> int | None:
    for name in names:
        if name.lower() in headers:
            return headers[name.lower()]
    if required:
        raise ValueError(
            f"En la hoja '{sheet}' falta la columna '{names[0]}' (encabezado en la primera fila)."
        )
    return None


def load_input(path: str) -> tuple[Config, list[Assignment], dict[str, set[Slot]]]:
    wb = load_workbook(path, data_only=True)

    if "Config" not in wb.sheetnames:
        raise ValueError("El archivo de entrada necesita una hoja llamada 'Config'.")
    ws = wb["Config"]
    headers = _header_map(ws)
    day_col = _col(headers, "dias", "días", "day", sheet="Config")
    period_col = _col(headers, "periodos", "períodos", "period", sheet="Config")

    days: list[str] = []
    periods: list[str] = []
    for row in ws.iter_rows(min_row=2):
        d = row[day_col - 1].value
        p = row[period_col - 1].value
        if d is not None and str(d).strip():
            days.append(str(d).strip())
        if p is not None and str(p).strip():
            periods.append(str(p).strip())
    if not days or not periods:
        raise ValueError("La hoja 'Config' debe tener al menos un día y un periodo.")

    config = Config(days=days, periods=periods)

    if "Asignaciones" not in wb.sheetnames:
        raise ValueError("El archivo de entrada necesita una hoja llamada 'Asignaciones'.")
    ws = wb["Asignaciones"]
    headers = _header_map(ws)
    group_col = _col(headers, "grupo", sheet="Asignaciones")
    subject_col = _col(headers, "materia", sheet="Asignaciones")
    teacher_col = _col(headers, "profesor", sheet="Asignaciones")
    hours_col = _col(headers, "horassemana", "horas semana", "horas", sheet="Asignaciones")
    room_col = _col(headers, "aula", "salon", "salón", required=False, sheet="Asignaciones")

    assignments: list[Assignment] = []
    for row in ws.iter_rows(min_row=2):
        group = row[group_col - 1].value
        if group is None or not str(group).strip():
            continue
        subject = row[subject_col - 1].value
        teacher = row[teacher_col - 1].value
        hours = row[hours_col - 1].value
        room = row[room_col - 1].value if room_col else None
        if subject is None or teacher is None or hours is None:
            raise ValueError(
                f"Fila incompleta en 'Asignaciones' para el grupo '{group}': "
                "faltan Materia, Profesor u HorasSemana."
            )
        hours = int(hours)
        if hours <= 0:
            continue
        assignments.append(
            Assignment(
                group=str(group).strip(),
                subject=str(subject).strip(),
                teacher=str(teacher).strip(),
                hours_per_week=hours,
                room=str(room).strip() if room not in (None, "") else None,
            )
        )
    if not assignments:
        raise ValueError("La hoja 'Asignaciones' no tiene ninguna fila con datos válidos.")

    unavailability: dict[str, set[Slot]] = {}
    if "NoDisponibilidad" in wb.sheetnames:
        ws = wb["NoDisponibilidad"]
        headers = _header_map(ws)
        teacher_col2 = _col(headers, "profesor", sheet="NoDisponibilidad")
        day_col2 = _col(headers, "dia", "día", sheet="NoDisponibilidad")
        period_col2 = _col(headers, "periodo", "período", sheet="NoDisponibilidad")
        for row in ws.iter_rows(min_row=2):
            teacher = row[teacher_col2 - 1].value
            day = row[day_col2 - 1].value
            period = row[period_col2 - 1].value
            if teacher is None or day is None or period is None:
                continue
            teacher = str(teacher).strip()
            unavailability.setdefault(teacher, set()).add(
                Slot(str(day).strip(), str(period).strip())
            )

    return config, assignments, unavailability


def _safe_sheet_name(name: str, used: set[str]) -> str:
    cleaned = re.sub(r"[:\\/?*\[\]]", "_", name).strip() or "Hoja"
    base = cleaned[:31]
    candidate = base
    n = 2
    while candidate in used:
        suffix = f"_{n}"
        candidate = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(candidate)
    return candidate


def _write_grid(
    ws: Worksheet,
    config: Config,
    title: str,
    cell_text: dict[tuple[str, str], str],
) -> None:
    ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=13)
    ws.cell(row=2, column=1, value="Periodo / Día")
    for col_i, day in enumerate(config.days, start=2):
        c = ws.cell(row=2, column=col_i, value=day)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = CELL_ALIGN

    for row_i, period in enumerate(config.periods, start=3):
        c = ws.cell(row=row_i, column=1, value=period)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = CELL_ALIGN
        for col_i, day in enumerate(config.days, start=2):
            text = cell_text.get((day, period), "")
            cell = ws.cell(row=row_i, column=col_i, value=text)
            cell.alignment = CELL_ALIGN

    ws.column_dimensions["A"].width = 16
    for col_i in range(2, 2 + len(config.days)):
        ws.column_dimensions[get_column_letter(col_i)].width = 22
    for row_i in range(3, 3 + len(config.periods)):
        ws.row_dimensions[row_i].height = 34


def write_output(
    path: str,
    config: Config,
    assignments: list[Assignment],
    schedule: dict[tuple[int, int], Slot],
) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    used_names: set[str] = set()

    groups = sorted({a.group for a in assignments})
    teachers = sorted({a.teacher for a in assignments})

    for group in groups:
        cell_text: dict[tuple[str, str], str] = {}
        for (a_idx, _occ), slot in schedule.items():
            a = assignments[a_idx]
            if a.group != group:
                continue
            label = a.subject + "\n" + a.teacher
            if a.room:
                label += "\n" + a.room
            cell_text[(slot.day, slot.period)] = label
        ws = wb.create_sheet(_safe_sheet_name(f"Grupo {group}", used_names))
        _write_grid(ws, config, f"Cronograma — Grupo {group}", cell_text)

    for teacher in teachers:
        cell_text = {}
        for (a_idx, _occ), slot in schedule.items():
            a = assignments[a_idx]
            if a.teacher != teacher:
                continue
            label = a.subject + "\n" + a.group
            if a.room:
                label += "\n" + a.room
            cell_text[(slot.day, slot.period)] = label
        ws = wb.create_sheet(_safe_sheet_name(f"Prof {teacher}", used_names))
        _write_grid(ws, config, f"Cronograma — Profesor {teacher}", cell_text)

    summary = wb.create_sheet(_safe_sheet_name("Resumen", used_names), 0)
    headers = ["Grupo", "Materia", "Profesor", "Aula", "Horas/semana", "Horarios asignados"]
    for col_i, h in enumerate(headers, start=1):
        c = summary.cell(row=1, column=col_i, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
    for row_i, a in enumerate(assignments, start=2):
        slots = sorted(
            (slot for (a_idx, _occ), slot in schedule.items() if a_idx == row_i - 2),
            key=lambda s: (config.days.index(s.day), config.periods.index(s.period)),
        )
        summary.cell(row=row_i, column=1, value=a.group)
        summary.cell(row=row_i, column=2, value=a.subject)
        summary.cell(row=row_i, column=3, value=a.teacher)
        summary.cell(row=row_i, column=4, value=a.room or "")
        summary.cell(row=row_i, column=5, value=a.hours_per_week)
        summary.cell(row=row_i, column=6, value=", ".join(str(s) for s in slots))
    for col_i, width in enumerate([14, 18, 18, 10, 12, 40], start=1):
        summary.column_dimensions[get_column_letter(col_i)].width = width

    wb.save(path)
